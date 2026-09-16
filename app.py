import io
import json
import base64
from google import genai
from google.genai import types
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from PIL import Image, ImageEnhance

st.set_page_config(
    page_title="学童野球スコア集計＆デジタル選手名鑑",
    page_icon="⚾️",
    layout="wide",
)

# ==========================================
# スマホ＆PC両立用レスポンシブCSSスタイル
# ==========================================
st.markdown("""
<style>
.stTabs [data-baseweb="tab-list"] {
    gap: 8px;
    overflow-x: auto !important;
    white-space: nowrap !important;
    padding-bottom: 6px;
    -webkit-overflow-scrolling: touch;
}
.stTabs [data-baseweb="tab"] {
    padding: 6px 14px;
    border-radius: 16px;
    background-color: rgba(120, 120, 120, 0.12);
    font-size: 0.9rem;
}
.sticky-mobile-viewer {
    position: -webkit-sticky;
    position: sticky;
    top: 3.5rem;
    z-index: 99;
    background-color: rgba(25, 25, 25, 0.95);
    padding: 8px;
    border-radius: 10px;
    box-shadow: 0 4px 15px rgba(0, 0, 0, 0.4);
    margin-bottom: 12px;
}
</style>
""", unsafe_allow_html=True)

# セッション状態の初期化
if "all_matches_data" not in st.session_state:
    st.session_state.all_matches_data = {}
if "match_images_b64" not in st.session_state:
    st.session_state.match_images_b64 = {}

# APIキー設定
api_key = st.secrets.get("GEMINI_API_KEY")
if not api_key:
    api_key = st.sidebar.text_input("管理者APIキー (Gemini)", type="password")

client = genai.Client(api_key=api_key) if api_key else None

# ==========================================
# 画像前処理：赤色インク強調画像の生成
# ==========================================
def enhance_red_pen(pil_img):
    """早稲田式の赤ペン結線（安打）を浮き彫りにする前処理画像を作成"""
    rgb_img = pil_img.convert("RGB")
    enhancer = ImageEnhance.Contrast(rgb_img)
    contrast_img = enhancer.enhance(1.4)
    
    r, g, b = contrast_img.split()
    enhanced_r = ImageEnhance.Brightness(r).enhance(1.2)
    merged = Image.merge("RGB", (enhanced_r, g, b))
    
    buf = io.BytesIO()
    merged.save(buf, format="JPEG", quality=90)
    return buf.getvalue()

# ==========================================
# 高精度化・改善版プロンプト群
# ==========================================

# Step 1: 選手枠特定用プロンプト（左端の選手列に完全集中）
ROSTER_PROMPT = """
### 役割
あなたは学童野球の手書きスコアブック（早稲田式）の「出場選手名簿抽出」の専門AIです。
スコアブックの左端にある「打順・背番号・選手名・守備位置」の列のみを精査し、出場選手の一覧を正確に抽出してください。

### 思考プロセス（必ずこの順で確認すること）
1. 打順1番から9番までを上から下へ走査する。
2. 各打順の枠内に「上下二段書き」があるかを確認する。
   - 上段：先発選手
   - 下段：途中交代選手、または代打選手（赤ペンでPH/代打の記載、あるいは別名が書かれている）
3. 背番号（1〜2桁の数字）と選手名（漢字・ひらがな）を正しくペアリングする。

### 立ち入り禁止区域（絶対に読み取らないこと）
- スコアシート最上部のイニング数字（1〜9）
- 右側のひし形マス目（打席結果）
- シート最下部の投球数、合計得点、審判名、メモ欄

### 制約条件
- [必須]: 上下二段書きがある場合、上段の選手と下段の選手を必ず別々のオブジェクトとして出力すること。
- [禁止]: 読み取れない文字を勝手に一般的な野球選手名などで推測補完すること。判読不能な部分は「未登録」とすること。

### 出力フォーマット（JSON配列のみを出力）
[
  {
    "batting_order": 1,
    "uniform_number": "10",
    "player_name": "山田",
    "is_substitute": false
  },
  {
    "batting_order": 1,
    "uniform_number": "15",
    "player_name": "田中",
    "is_substitute": true
  }
]
"""

# Step 2: 打席・イニング・打点・盗塁判定用プロンプト（四角結線×得点丸判定 ＆ 白旗「要確認」ルール）
DETAILS_PROMPT = """
### 役割
あなたは学童野球の手書きスコアブック（早稲田式）の「打撃結果・安打結線」を高精度に判定する専門AIです。
事前に確定された【選手リスト】に基づき、スコアブック画像および赤ペン強調画像から、各選手の1回〜7回の打席結果、打点、盗塁数を判定してください。

### 【絶対厳守】迷ったら推測せず「要確認」にする白旗ルール
- 本チームのスコアブックは文字による「HR」等のメモがなく、純粋に「赤線のみ」で安打種別を記録しています。
- そのため、保護者による手書きのインク太さ・滲み・重なりによって「2塁打などの安打線＋進塁線＋中央の得点丸数字」が混ざり、本塁打なのか2塁打や単打の生還なのか100%確信が持てないケースがあります。
- このような曖昧さや迷いがある場合は、絶対に勝手な推測で本塁打や単打と決めつけてはいけません。
- 少しでも判定に自信がない・線の本数や進塁線との重なりが怪しい打席は、必ず結果を `"要確認"` としてください。人間が後から画面でピンポイント修正します。

### サボり・読み飛ばし防止ルール（全マス完全走査の義務付け）
- 対象となる選手全員について、必ず「第1回」から「第7回」までの全7マスを1マスも省略せずに順番にチェックしてください。
- 打席が回っていないイニングは、適当に前の回をコピーしたり省略したりせず、必ず `"なし"` と明記してください。
- 途中の回や下位打線で集中力を切らして推測で埋める行為は厳禁です。最低でも21箇所以上の全打席マスについて、以下の基準で1マスずつ厳密に判定してください。

### イニング列ズレ防止ルール（野球の連続性検証）
1. スコアシート最上部の大きな数字（1〜7）の真下にあるマス目のみをそのイニングの打席とします。
2. 【打順の連続性チェック】:
   - 前のイニングの最後（3アウトチェンジ＝斜め二重線「//」や丸数字③）が「3番打者」だった場合、次のイニングの先頭は必ず「4番打者」から始まります。
   - もし次の回の打席が1番打者の位置に見えた場合、それは列の誤認です。打順の連続性とイニング終了線（//）をアンカーにして、空欄イニングを絶対に詰めずに正しく「なし」と判定してください。

### 「四角（◇）結線 × 中の丸数字（得点）」の厳密判定基準
中央のひし形（ダイヤモンド）の「赤ペンの結線」と「得点丸数字」を確認してください：

1. **【本塁打（ホームラン）の条件】**:
   - ダイヤモンド外枠の4辺（本塁→一塁→二塁→三塁→本塁）すべてが打者自身の赤線で完全に一周囲まれている（◇の四角形が完成している）。
   - かつ、その中央に「丸の中に数字（得点丸数字）」が書かれている。
   - ※外枠の◇が4辺すべて赤線で閉じている確信がない限り、中央に丸数字があっても絶対に本塁打にしてはいけません。迷う場合は `"要確認"` としてください。

2. **【単打・四球・凡打 ＋ 得点（ホームイン）の条件】**:
   - 中央に「丸の中に数字（得点丸数字）」があるが、外枠の◇が赤線で一周していない場合。
   - 例：右下の辺（本塁→一塁）のみ赤線で、中央に丸数字 ➜ 打撃結果は【単打】（後続打者の安打等で本塁生還した記録）。
   - 例：外枠に赤線がなく鉛筆で「四」や「B」、中央に丸数字 ➜ 打撃結果は【四球】。
   - ※中央の得点丸に惑わされて、単打や四球を本塁打と誤認してはいけません。

3. **【2塁打・3塁打の条件】**:
   - 2塁打: 本塁→一塁→二塁（真上の頂点）まで連続した赤線がある（外枠が半周）。
   - 3塁打: 本塁→一塁→二塁→三塁（左端頂点）まで連続した赤線がある。
   - ※2塁打を打って後から生還したため四角が一周して見える場合で、本塁打か2塁打か判別が難しい場合は `"要確認"` としてください。

4. **【凡打・四死球・犠打】**:
   - 凡打: 外枠に赤線がなく、黒字で「3A」「4-3」「K」「8」などのアウト記号、または丸数字（①②③）がある。
   - 四球/死球: 「四」「B」または「死」「DB」。
   - 犠打: 「□」「サ」などの犠打記号。

### 盗塁（stolen_bases）と打点（rbi）の判定
- **盗塁 (stolen_bases)**: ダイヤモンドの進塁線付近に「S」または「ㄭ」の記号がある個数をカウント（盗塁進塁線を安打の赤線と混同しないこと）。
- **打点 (rbi)**: 適時打のマス目隅に小さく書かれた数字、または明確なタイムリー打席から計算（不明・なしなら0）。

### 括弧書き数字「(数字)」の除外（最重要）
- マス目内の「(6)」「(8)」などの括弧付き数字は「進塁責任打者」の記録です。打者本人の打撃結果や打点コードではありません。完全に無視してください。

### 立ち入り禁止区域（除外エリア）
- 各イニング列の最下段にある「得点・残塁・投球数」の集計行
- シート右端の「通算打数・安打」集計欄
- 欄外のメモや審判欄

### 出力フォーマット（指定のJSON配列のみ返却）
※inningsの値には必ず "単打", "2塁打", "3塁打", "本塁打", "四球", "死球", "三振", "凡打", "犠打", "要確認", "なし" のいずれかを使用してください。
[
  {
    "match_date": "試合日（不明なら -）",
    "opponent": "対戦相手（不明なら -）",
    "batting_order": 1,
    "uniform_number": "10",
    "player_name": "山田",
    "is_substitute": false,
    "rbi": 1,
    "stolen_bases": 1,
    "innings": {
      "1": "単打",
      "2": "なし",
      "3": "要確認",
      "4": "なし",
      "5": "なし",
      "6": "なし",
      "7": "なし"
    },
    "highlight": "初回先頭打者安打と盗塁"
  }
]
"""

# 「要確認」を選択肢に追加
RESULT_OPTIONS = ["なし", "要確認", "単打", "2塁打", "3塁打", "本塁打", "四球", "死球", "三振", "凡打", "犠打"]

def calculate_stats_from_grid(grid_players, match_file_name=""):
    compiled_records = []
    for p in grid_players:
        ab = 0
        hits = 0
        doubles = 0
        triples = 0
        hrs = 0
        so = 0
        bb = 0
        db = 0
        pa = 0

        for inn_str in ["1", "2", "3", "4", "5", "6", "7"]:
            res = p["innings"].get(inn_str, "なし")
            if res == "なし":
                continue
            pa += 1
            if res == "単打":
                hits += 1
                ab += 1
            elif res == "2塁打":
                doubles += 1
                ab += 1
            elif res == "3塁打":
                triples += 1
                ab += 1
            elif res == "本塁打":
                hrs += 1
                ab += 1
            elif res == "三振":
                so += 1
                ab += 1
            elif res == "凡打":
                ab += 1
            elif res == "四球":
                bb += 1
            elif res == "死球":
                db += 1
            elif res == "犠打":
                pass
            elif res == "要確認":
                # 未確定のまま残っていた場合は打席数のみカウントし、安打や凡打には加算しない安全処理
                pass

        compiled_records.append({
            "source_file": match_file_name,
            "match_date": p.get("match_date", "-"),
            "opponent": p.get("opponent", "-"),
            "batting_order": p.get("batting_order", 0),
            "uniform_number": str(p.get("uniform_number", "")).strip(),
            "player_name": str(p.get("player_name", "未登録")).strip(),
            "plate_appearances": pa,
            "at_bats": ab,
            "hits": hits,
            "doubles": doubles,
            "triples": triples,
            "homeruns": hrs,
            "strikeouts": so,
            "walks": bb,
            "dead_ball": db,
            "stolen_bases": int(p.get("stolen_bases", 0)),
            "rbi": int(p.get("rbi", 0)),
            "runs": 0,
            "highlight": p.get("highlight", "")
        })
    return compiled_records

def create_excel_from_compiled(all_records):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    players = {}
    for r in all_records:
        name = (r.get("player_name") or "未登録").strip()
        players.setdefault(name, []).append(r)

    headers = [
        "背番号", "試合日", "対戦相手", "打席", "打数", "安打", "2塁打", "3塁打", "本塁打",
        "三振", "四球", "死球", "盗塁", "打点", "得点", "ハイライト"
    ]

    for player_name, matches in players.items():
        sheet_title = player_name[:30] if player_name else "未登録"
        ws = wb.create_sheet(title=sheet_title)
        ws.append(headers)

        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")

        for m in matches:
            ws.append([
                m.get("uniform_number", ""),
                m.get("match_date", "-"),
                m.get("opponent", "-"),
                m.get("plate_appearances", 0),
                m.get("at_bats", 0),
                m.get("hits", 0),
                m.get("doubles", 0),
                m.get("triples", 0),
                m.get("homeruns", 0),
                m.get("strikeouts", 0),
                m.get("walks", 0),
                m.get("dead_ball", 0),
                m.get("stolen_bases", 0),
                m.get("rbi", 0),
                m.get("runs", 0),
                m.get("highlight", "")
            ])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

# タブ構成
tab_admin, tab_kids = st.tabs(
    ["📝 役員用（複数試合一括解析＆ポチポチ確定）", "🏆 選手名鑑＆アワード"]
)

# ==========================================
# ① 役員用（2段階高精度解析 ＆ 照合エディタ）
# ==========================================
with tab_admin:
    st.subheader("手書きスコア解析 ＆ 照合エディタ（高精度サボり防止エンジン）")
    st.caption("赤線強調フィルタ、打順連続性チェック、四角結線×得点丸の厳格分離により、全打席を高精度に走査・判定します。")

    if not client:
        st.warning("Gemini APIキーを設定してください（Secrets または サイドバー）。")
        st.stop()

    uploaded_files = st.file_uploader(
        "スコアブック写真を選択（複数ファイル選択可）",
        type=["jpg", "jpeg", "png"],
        accept_multiple_files=True
    )

    if uploaded_files:
        if st.button(f"AIで全{len(uploaded_files)}試合を高精度一括解析する", type="primary"):
            progress_bar = st.progress(0)
            status_text = st.empty()
            
            new_all_matches = {}
            new_images_b64 = {}

            for idx, f in enumerate(uploaded_files):
                f_name = f.name
                status_text.text(f"【{idx+1}/{len(uploaded_files)}】{f_name} の赤ペン強調＆選手名簿を確定中...")
                raw_bytes = f.read()
                new_images_b64[f_name] = base64.b64encode(raw_bytes).decode()
                
                # 画像の赤ペン強調処理
                pil_img = Image.open(io.BytesIO(raw_bytes))
                enhanced_red_bytes = enhance_red_pen(pil_img)

                try:
                    # Step 1: 選手名簿の確定
                    res_roster = client.models.generate_content(
                        model="gemini-3.6-flash",
                        contents=[
                            types.Part.from_bytes(data=raw_bytes, mime_type="image/jpeg"),
                            "スコアブック左側の打順・背番号・選手名（先発・交代・代打二段書き含む）を漏れなく抽出してください。"
                        ],
                        config=types.GenerateContentConfig(
                            system_instruction=ROSTER_PROMPT,
                            response_mime_type="application/json",
                            temperature=0.1
                        )
                    )
                    roster_data = res_roster.text

                    # Step 2: 選手枠に基づき全打席マス目・安打結線・得点丸を完全走査
                    status_text.text(f"【{idx+1}/{len(uploaded_files)}】{f_name} の全イニング打席を走査中（迷ったら「要確認」）...")
                    res_details = client.models.generate_content(
                        model="gemini-3.6-flash",
                        contents=[
                            types.Part.from_bytes(data=raw_bytes, mime_type="image/jpeg"),
                            types.Part.from_bytes(data=enhanced_red_bytes, mime_type="image/jpeg"),
                            f"確定選手名簿:\n{roster_data}\n\n上記選手枠に基づき、スコアブックの1回〜7回の全打席詳細、打点、盗塁を判定してください。赤線と得点丸が重なって判別できない打席は迷わず「要確認」としてください。"
                        ],
                        config=types.GenerateContentConfig(
                            system_instruction=DETAILS_PROMPT,
                            response_mime_type="application/json",
                            temperature=0.1
                        )
                    )
                    parsed = json.loads(res_details.text)
                    new_all_matches[f_name] = parsed

                except Exception as e:
                    st.error(f"{f_name} の解析エラー: {e}")

                progress_bar.progress((idx + 1) / len(uploaded_files))

            status_text.empty()
            if new_all_matches:
                st.session_state.all_matches_data = new_all_matches
                st.session_state.match_images_b64 = new_images_b64
                st.success(f"🎉 全 {len(new_all_matches)} 試合分の解析が完了しました！")

    # 照合・編集盤面
    if st.session_state.all_matches_data:
        st.divider()
        match_files = list(st.session_state.all_matches_data.keys())
        
        top_c1, top_c2 = st.columns([2, 1])
        selected_match_file = top_c1.selectbox("📁 確認・編集する試合を選択", match_files)
        is_mobile_sticky = top_c2.checkbox("📱 スマホ表示（画像を上部に固定）", value=False)

        current_players = st.session_state.all_matches_data.get(selected_match_file, [])
        current_b64 = st.session_state.match_images_b64.get(selected_match_file, "")

        add_col1, add_col2 = st.columns([1, 3])
        with add_col1:
            if st.button("➕ この試合に選手を手動追加"):
                new_player_template = {
                    "batting_order": len(current_players) + 1,
                    "uniform_number": "",
                    "player_name": f"追加選手{len(current_players) + 1}",
                    "is_substitute": True,
                    "rbi": 0,
                    "stolen_bases": 0,
                    "innings": {"1": "なし", "2": "なし", "3": "なし", "4": "なし", "5": "なし", "6": "なし", "7": "なし"},
                    "highlight": ""
                }
                st.session_state.all_matches_data[selected_match_file].append(new_player_template)
                st.rerun()

        col_img, col_grid = st.columns([1.1, 1.3])

        # 左側：画像ビューワー
        with col_img:
            st.markdown(f"#### 📷 原本画像: `{selected_match_file}`")
            zoom_val = st.slider("🔍 拡大率", min_value=100, max_value=350, value=150, step=25, format="%d%%")
            
            box_height = 320 if is_mobile_sticky else 620
            sticky_class = "sticky-mobile-viewer" if is_mobile_sticky else ""

            viewer_html = f"""
            <div class="{sticky_class}" style="width:100%; height:{box_height}px; overflow:auto; border:2px solid #555; border-radius:8px; background-color:#222; text-align:center;">
                <img src="data:image/jpeg;base64,{current_b64}" style="width:{zoom_val}%; max-width:none; transition:width 0.15s ease-in-out; cursor:grab;" />
            </div>
            """
            components.html(viewer_html, height=box_height + 20)

        # 右側：付箋タブエディタ
        with col_grid:
            st.markdown("#### 🎯 打席盤面エディタ")
            st.caption("タブを指で横にスワイプして選手を選択できます。「要確認」の打席を中心に修正してください。")

            tab_labels = []
            for idx, player in enumerate(current_players):
                u_num = str(player.get("uniform_number", "")).strip()
                num_str = f"#{u_num} " if u_num else ""
                p_name = str(player.get("player_name", "選手")).strip()
                sub_tag = "(代)" if player.get("is_substitute") else ""
                tab_labels.append(f"{num_str}{p_name}{sub_tag}")

            player_tabs = st.tabs(tab_labels)
            edited_current_players = []

            for idx, (p_tab, player) in enumerate(zip(player_tabs, current_players)):
                with p_tab:
                    is_sub = player.get("is_substitute", False)
                    order_val = player.get("batting_order", idx + 1)

                    st.markdown(f"##### **選手情報設定 {'（途中交代・代打）' if is_sub else '（先発）'}**")

                    p_cols = st.columns([1, 2, 3])
                    u_num = p_cols[0].text_input("背番号", value=str(player.get("uniform_number", "")), key=f"{selected_match_file}_num_{idx}")
                    p_name = p_cols[1].text_input("選手名（漢字）", value=str(player.get("player_name", "")), key=f"{selected_match_file}_name_{idx}")
                    hl = p_cols[2].text_input("ハイライトメモ", value=str(player.get("highlight", "")), key=f"{selected_match_file}_hl_{idx}")

                    stat_c1, stat_c2 = st.columns(2)
                    rbi_val = stat_c1.number_input("打点 (RBI)", min_value=0, max_value=20, value=int(player.get("rbi", 0)), step=1, key=f"{selected_match_file}_rbi_{idx}")
                    sb_val = stat_c2.number_input("盗塁数 (SB)", min_value=0, max_value=20, value=int(player.get("stolen_bases", 0)), step=1, key=f"{selected_match_file}_sb_{idx}")

                    st.markdown("**各イニングの打撃結果（1回〜7回）**")
                    inn_cols = st.columns(7)
                    new_innings = {}
                    for i_idx, inn_str in enumerate(["1", "2", "3", "4", "5", "6", "7"]):
                        cur_val = player.get("innings", {}).get(inn_str, "なし")
                        default_idx = RESULT_OPTIONS.index(cur_val) if cur_val in RESULT_OPTIONS else 0
                        sel = inn_cols[i_idx].selectbox(
                            f"{inn_str}回",
                            RESULT_OPTIONS,
                            index=default_idx,
                            key=f"{selected_match_file}_inn_{idx}_{inn_str}"
                        )
                        new_innings[inn_str] = sel

                    edited_current_players.append({
                        "match_date": player.get("match_date", "-"),
                        "opponent": player.get("opponent", "-"),
                        "batting_order": order_val,
                        "uniform_number": u_num,
                        "player_name": p_name,
                        "is_substitute": is_sub,
                        "rbi": rbi_val,
                        "stolen_bases": sb_val,
                        "innings": new_innings,
                        "highlight": hl
                    })

            st.session_state.all_matches_data[selected_match_file] = edited_current_players

            st.write("")
            if st.button("💾 全試合の成績を統合確定・Excelを作成する", type="primary", use_container_width=True):
                all_compiled = []
                for m_file, p_list in st.session_state.all_matches_data.items():
                    compiled_single = calculate_stats_from_grid(p_list, match_file_name=m_file)
                    all_compiled.extend(compiled_single)

                st.session_state["compiled_records"] = all_compiled
                st.success(f"🎉 全 {len(st.session_state.all_matches_data)} 試合分の成績を確定統合しました！")

        if "compiled_records" in st.session_state:
            st.divider()
            excel_data = create_excel_from_compiled(st.session_state["compiled_records"])
            st.download_button(
                label=f"📥 全{len(st.session_state.all_matches_data)}試合分 選手名別シート付きExcelをダウンロード",
                data=excel_data,
                file_name="チーム通算打撃成績一覧.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

# ==========================================
# ② 選手名鑑＆アワード（全試合通算）
# ==========================================
with tab_kids:
    compiled_data = st.session_state.get("compiled_records")
    if not compiled_data:
        st.info("👈 まず「役員用」タブでスコアを確定させてください。")
    else:
        df = pd.DataFrame(compiled_data)
        players = [p for p in df["player_name"].dropna().unique().tolist() if p.strip() != ""]

        st.subheader("🎖️ チームタイトル・アワード（通算集計）")
        c1, c2, c3, c4 = st.columns(4)

        total_h = df["hits"] + df["doubles"] + df["triples"] + df["homeruns"]
        df["total_hits"] = total_h
        tb_leaders = df.groupby("player_name")["total_hits"].sum().sort_values(ascending=False)
        if not tb_leaders.empty and tb_leaders.iloc[0] > 0:
            c1.metric("最多安打（長打含む）", f"{tb_leaders.index[0]} 選手", f"{int(tb_leaders.iloc[0])} 本")

        hr_leaders = df.groupby("player_name")["homeruns"].sum().sort_values(ascending=False)
        if not hr_leaders.empty and hr_leaders.iloc[0] > 0:
            c2.metric("スラッガー賞（本塁打）", f"{hr_leaders.index[0]} 選手", f"{int(hr_leaders.iloc[0])} 本")

        rbi_leaders = df.groupby("player_name")["rbi"].sum().sort_values(ascending=False)
        if not rbi_leaders.empty and rbi_leaders.iloc[0] > 0:
            c3.metric("クラッチヒッター賞（打点）", f"{rbi_leaders.index[0]} 選手", f"{int(rbi_leaders.iloc[0])} 打点")

        sb_leaders = df.groupby("player_name")["stolen_bases"].sum().sort_values(ascending=False)
        if not sb_leaders.empty and sb_leaders.iloc[0] > 0:
            c4.metric("スピードスター賞（盗塁）", f"{sb_leaders.index[0]} 選手", f"{int(sb_leaders.iloc[0])} 個")

        st.divider()
        st.subheader("⚾️ チーム デジタル選手名鑑（全試合通算）")
        selected_player = st.selectbox("選手を選択してください（名前で通算集計）", players)
        player_data = df[df["player_name"] == selected_player]

        used_numbers = [str(n).strip() for n in player_data["uniform_number"].dropna().unique() if str(n).strip() != ""]
        num_display = f"（着用背番号: #{', #'.join(used_numbers)}）" if used_numbers else ""

        ab = player_data["at_bats"].sum()
        h = player_data["total_hits"].sum()
        avg = (h / ab) if ab > 0 else 0.0
        hr = player_data["homeruns"].sum()
        rbi = player_data["rbi"].sum()
        sb = player_data["stolen_bases"].sum()
        matches_count = player_data["source_file"].nunique()

        st.markdown(f"### **{selected_player}** 選手の確定通算成績 {num_display}")
        st.caption(f"集計対象: 全 {matches_count} 試合出場")
        
        col1, col2, col3, col4, col5 = st.columns(5)
        col1.metric("通算打率", f".{int(avg * 1000):03d}" if avg > 0 else ".000")
        col2.metric("通算安打", f"{int(h)} 本")
        col3.metric("本塁打", f"{int(hr)} 本")
        col4.metric("通算打点", f"{int(rbi)} 点")
        col5.metric("通算盗塁", f"{int(sb)} 個")

        st.markdown("#### 🔥 各試合のベストハイライト")
        hl_list = player_data[player_data["highlight"].str.strip() != ""]["highlight"].tolist()
        if hl_list:
            for hl in hl_list:
                st.write(f"・{hl}")
        else:
            st.write("・全力プレーでチームに大きく貢献！")
